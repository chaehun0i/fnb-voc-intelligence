"""생성한 제품 Template과 제한된 파일 parser. 외부 연결/수식 실행 없음."""
import csv
from io import BytesIO, StringIO
from pathlib import PurePosixPath
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from src.application.intake_schema import SCHEMAS

MAX_BYTES, MAX_ROWS, MAX_COLUMNS, MAX_SHEETS = 2_000_000, 1000, 20, 5


class IntakeFileError(ValueError):
    pass


def parse_file(filename, content, content_type, csv_kind="VOC"):
    """읽기 전에 크기/ZIP/차원을 제한하고 원본 파일은 보관하지 않습니다."""
    if not filename or len(filename) > 200 or "/" in filename or "\\" in filename or "\x00" in filename:
        raise IntakeFileError("파일명에 경로를 사용할 수 없습니다.")
    extension = PurePosixPath(filename).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise IntakeFileError(".xlsx 또는 UTF-8 .csv 파일만 지원합니다.")
    if not content or len(content) > MAX_BYTES:
        raise IntakeFileError("빈 파일이거나 최대 크기 2MB를 초과했습니다.")
    try:
        if extension == ".csv":
            if content.startswith(b"PK") or content_type not in {"text/csv", "text/plain", "application/csv", "application/octet-stream", "application/vnd.ms-excel"}:
                raise IntakeFileError("CSV 파일 형식을 확인해 주세요.")
            if csv_kind not in SCHEMAS:
                raise IntakeFileError("CSV 데이터 종류를 선택해 주세요.")
            reader = csv.reader(StringIO(content.decode("utf-8-sig")), strict=True)
            rows = []
            for row in reader:
                if len(rows) > MAX_ROWS or len(row) > MAX_COLUMNS or any(len(c) > 4000 for c in row):
                    raise IntakeFileError("최대 1000행·20열, 셀 4000자까지 지원합니다.")
                rows.append(row)
            tables = [(csv_kind, rows)]
        else:
            if not content.startswith(b"PK\x03\x04") or content_type not in {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/octet-stream"}:
                raise IntakeFileError("XLSX signature/MIME를 확인해 주세요.")
            with ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 100 or sum(e.file_size for e in entries) > 20_000_000:
                    raise IntakeFileError("압축 해제 크기 또는 내부 파일 수가 한도를 초과했습니다.")
                for entry in entries:
                    path = entry.filename.lower()
                    if ".." in PurePosixPath(path).parts or path.startswith("/") or any(x in path for x in ("vbaproject", "externallinks", "embeddings/", "activex/")):
                        raise IntakeFileError("매크로·외부 링크·삽입 개체는 지원하지 않습니다.")
                    if path.endswith((".xml", ".rels")):
                        xml = archive.read(entry)
                        if b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper() or b"macroenabled" in xml.lower():
                            raise IntakeFileError("XML 외부 선언은 허용하지 않습니다.")
            book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.worksheets) > MAX_SHEETS:
                    raise IntakeFileError("최대 5개 Sheet를 지원합니다.")
                tables = []
                for sheet in book.worksheets:
                    if (sheet.max_row or 0) > MAX_ROWS+1 or (sheet.max_column or 0) > MAX_COLUMNS:
                        raise IntakeFileError("최대 1000행·20열을 지원합니다.")
                    rows = []
                    for cells in sheet.iter_rows():
                        if len(rows) > MAX_ROWS or len(cells) > MAX_COLUMNS:
                            raise IntakeFileError("실제 입력은 최대 1000행·20열까지 지원합니다.")
                        if any(c.data_type == "f" for c in cells):
                            raise IntakeFileError("Excel 수식은 지원하지 않습니다. 값으로 붙여 넣어 주세요.")
                        row = [c.value if c.value is not None else "" for c in cells]
                        if any(len(str(v)) > 4000 for v in row):
                            raise IntakeFileError("셀은 최대 4000자까지 지원합니다.")
                        rows.append(row)
                    if sheet.title != "입력가이드":
                        tables.append((sheet.title, rows))
            finally:
                book.close()
        if not tables or sum(max(0, len(rows)-1) for _, rows in tables) > MAX_ROWS:
            raise IntakeFileError("최대 합계 1000행까지 지원합니다.")
        for _, rows in tables:
            if not rows or not any(any(str(v).strip() for v in row) for row in rows[1:]):
                raise IntakeFileError("Header와 입력 행이 필요합니다.")
            headers = [str(v).strip() for v in rows[0]]
            if any(not h or len(h) > 80 for h in headers) or len(set(headers)) != len(headers):
                raise IntakeFileError("빈 Header 또는 중복 Header를 수정해 주세요.")
            rows[0] = headers
        return tables
    except IntakeFileError:
        raise
    except (BadZipFile, UnicodeError, ValueError, KeyError, TypeError, csv.Error, OSError,
            ParseError, DefusedXmlException, InvalidFileException, IndexError, RuntimeError):
        raise IntakeFileError("파일을 읽을 수 없습니다. 손상 여부와 UTF-8 형식을 확인해 주세요.") from None


def template_bytes():
    book = Workbook()
    guide = book.active
    guide.title = "입력가이드"
    guide.append(["ServIQ 입력 템플릿 v1", "개인정보는 입력하지 마세요."])
    guide.append(["입력 방법", "매장명을 UI에서 먼저 등록하고 같은 이름을 모든 행에 사용하세요."])
    guide.append(["자료ID", "같은 원본 자료에는 같은 ID. 영문/숫자/._:-, 최대 64자."])
    guide.append(["발생일시", "ISO 8601, 시간대 포함. 날짜만 있으면 UTC 00:00으로 처리합니다."])
    guide.append(["판매_거래 / 재고", "Read-only Observation 입력입니다. 실제 POS/ERP Connector나 매출·재고 원장 동기화가 아닙니다."])
    guide.append(["거래 관측코드", "REFUND_SIGNAL / CANCEL_SIGNAL"])
    guide.append(["재고 관측코드", "STOCK_SHORTAGE / STOCK_ADJUSTMENT"])
    guide.append(["주의", "예제 행을 운영 데이터로 교체하세요. 수식/매크로/외부 링크는 허용하지 않습니다."])
    for name, columns in SCHEMAS.items():
        sheet = book.create_sheet(name)
        sheet.append([c.label for c in columns])
        sheet.append([c.example for c in columns])
        guide.append([name, ", ".join(f"{c.label}: {c.type} ({'필수' if c.required else '선택'})" for c in columns)])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="234B63")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 34 if column[0].value != "VOC 내용" else 65
    guide.column_dimensions["A"].width = 28
    guide.column_dimensions["B"].width = 115
    output = BytesIO()
    book.save(output)
    return output.getvalue()
